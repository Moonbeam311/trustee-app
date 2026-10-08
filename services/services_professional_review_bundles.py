"""Native, non-execution Professional Review bundle generation.

The existing ``generated_documents`` table owns the persisted bundle identity.
Components are derived files; this module creates no registry, review, or
provenance schema and never changes a source or Professional Review record.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import sqlite3
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document


WILL_BUNDLE = "WILL_PROFESSIONAL_REVIEW"
TR001_BUNDLE = "TR001_NJ_PROFESSIONAL_REVIEW"
WILL_STATUS = "PROFESSIONAL REVIEW CANDIDATE — NOT FOR SIGNATURE OR EXECUTION"
TR001_STATUS = (
    "PROFESSIONAL REVIEW MATERIAL — NOT A TRUST CREATION, EXECUTION, AMENDMENT, "
    "ACCEPTANCE, OR FUNDING INSTRUMENT"
)
TR001_TRUST_STATUS = "CONTROLLED PRE-SIGNING — EXECUTION NOT YET COMPLETED"

PROFILES = {
    WILL_BUNDLE: {
        "template_id": "REVIEW-BUNDLE-WILL-V1",
        "source_type": "intake",
        "status": WILL_STATUS,
        "components": (
            "Integrated_Will_Professional_Review_Candidate.docx",
            "Professional_Review_Issues_Memorandum.docx",
            "Architecture_and_Decision_Crosswalk.docx",
            "Candidate_Control_Manifest.json",
        ),
    },
    TR001_BUNDLE: {
        "template_id": "REVIEW-BUNDLE-TR001-NJ-V1",
        "source_type": "trust",
        "status": TR001_STATUS,
        "components": (
            "New_Jersey_Counsel_Review_Packet.docx",
            "Tax_Professional_Review_Packet.docx",
            "Evidence_and_Provenance_Index.docx",
            "Review_Control_Manifest.json",
        ),
    },
}


class ReviewBundleError(ValueError):
    pass


def _json(value: Any) -> str:
    return json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":"), default=str)


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _tables(connection: sqlite3.Connection) -> set[str]:
    return {row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table'")}


def _columns(connection: sqlite3.Connection, table: str) -> set[str]:
    return {row[1] for row in connection.execute(f"PRAGMA table_info({table})")}


def _one(connection: sqlite3.Connection, table: str, key: str, value: str, firm_id: str) -> dict[str, Any]:
    if table not in _tables(connection) or key not in _columns(connection, table):
        raise ReviewBundleError(f"canonical_{table}_owner_unavailable")
    clauses, params = [f"{key}=?"], [value]
    if "firm_id" in _columns(connection, table):
        clauses.append("firm_id=?"); params.append(firm_id)
    connection.row_factory = sqlite3.Row
    row = connection.execute(f"SELECT * FROM {table} WHERE {' AND '.join(clauses)}", params).fetchone()
    if not row:
        raise ReviewBundleError(f"canonical_{table}_record_unavailable")
    return dict(row)


def _canonical_source(connection, bundle_type, source_id, firm_id):
    if bundle_type == TR001_BUNDLE:
        return _one(connection, "trusts", "trust_id", source_id, firm_id)
    candidates = (("intake_sessions", "intake_id"), ("intakes", "intake_id"))
    for table, key in candidates:
        if table in _tables(connection) and key in _columns(connection, table):
            return _one(connection, table, key, source_id, firm_id)
    raise ReviewBundleError("canonical_intake_owner_unavailable")


def _issues(connection, intake_id, firm_id):
    if not intake_id or "professional_review_issues" not in _tables(connection):
        return []
    connection.row_factory = sqlite3.Row
    return [dict(row) for row in connection.execute(
        "SELECT * FROM professional_review_issues WHERE intake_id=? AND firm_id=? ORDER BY issue_id",
        (intake_id, firm_id),
    )]


_RESOLVED_REVIEW_STATES = {
    "resolved", "closed", "closed_no_conflict", "superseded", "non_applicable",
    "not_applicable", "accepted_risk",
}


def _active_review_issue(row):
    status = str(row.get("status") or "").strip().lower().replace("-", "_").replace(" ", "_")
    disposition = str(row.get("disposition") or "").strip().lower().replace("-", "_").replace(" ", "_")
    return status not in _RESOLVED_REVIEW_STATES and disposition not in _RESOLVED_REVIEW_STATES


def _unresolved_item(*, identity, subject, reason, owner, linked_type=None,
                     linked_id=None, bundle_type=None, context_id=None):
    return {
        "blocker_id": identity, "subject": subject,
        "state": "INTENTIONAL_UNRESOLVED", "reason": reason,
        "source_owner": owner, "source_type": owner,
        "linked_record_type": linked_type, "linked_record_id": linked_id,
        "affected_bundle": bundle_type, "affected_context_id": context_id,
    }


def _derive_unresolved_items(connection, *, bundle_type, source, source_id, issues):
    items = []
    for row in issues:
        if _active_review_issue(row):
            items.append(_unresolved_item(
                identity=row["issue_id"], subject=row.get("reviewer_notes") or row["issue_id"],
                reason="Canonical Professional Review issue remains open or review-needed",
                owner="professional_review_issues",
                linked_type=row.get("linked_record_type"), linked_id=row.get("linked_record_id"),
                bundle_type=bundle_type, context_id=source_id,
            ))
    if bundle_type == TR001_BUNDLE:
        field_rules = {
            "execution_status": ("execution", {"executed", "complete", "completed", "finalized"}),
            "funding_status": ("funding", {"funded", "complete", "completed"}),
        }
        for field, (subject, resolved) in field_rules.items():
            if field in source and source.get(field) not in (None, ""):
                value = str(source[field]).strip()
                if value.lower().replace("-", "_").replace(" ", "_") not in resolved:
                    items.append(_unresolved_item(
                        identity=f"trusts:{source_id}:{field}", subject=subject,
                        reason=f"Canonical Trust {field} is {value}", owner="trusts",
                        linked_type="trust", linked_id=source_id,
                        bundle_type=bundle_type, context_id=source_id,
                    ))
        tables = _tables(connection)
        if "successor_acceptances" in tables:
            columns = _columns(connection, "successor_acceptances")
            if {"acceptance_id", "trust_id", "firm_id", "acceptance_status"}.issubset(columns):
                rows = [dict(row) for row in connection.execute(
                    "SELECT * FROM successor_acceptances WHERE trust_id=? AND firm_id=?",
                    (source_id, source.get("firm_id")),
                )]
                prior_key = "prior_acceptance_id" if "prior_acceptance_id" in columns else None
                terminals = (_lineage_terminals(rows, "acceptance_id", prior_key)
                             if prior_key else rows)
                for row in terminals:
                    state = str(row.get("acceptance_status") or "").strip()
                    if state.upper() in {"PENDING_EVIDENCE", "PENDING", "REVIEW_REQUIRED", "UNRESOLVED"}:
                        items.append(_unresolved_item(
                            identity=row["acceptance_id"], subject="trustee_acceptance",
                            reason=f"Canonical acceptance status is {state}",
                            owner="successor_acceptances", linked_type="trust",
                            linked_id=source_id, bundle_type=bundle_type,
                            context_id=source_id,
                        ))
        if "hub_authority_applicability" in tables:
            columns = _columns(connection, "hub_authority_applicability")
            needed = {"applicability_id", "firm_id", "context_type", "context_id"}
            if needed.issubset(columns):
                rows = [dict(row) for row in connection.execute(
                    "SELECT * FROM hub_authority_applicability WHERE firm_id=? AND context_type='TRUST' AND context_id=?",
                    (source.get("firm_id"), source_id),
                )]
                prior_key = "prior_applicability_id" if "prior_applicability_id" in columns else None
                terminals = (_lineage_terminals(rows, "applicability_id", prior_key)
                             if prior_key else rows)
                for row in terminals:
                    authorized = row.get("generation_authorized")
                    support = str(row.get("independent_support_state") or "").upper()
                    if authorized in (0, "0", False) or (support and support != "YES"):
                        items.append(_unresolved_item(
                            identity=row["applicability_id"],
                            subject=row.get("subject") or "jurisdiction_applicability",
                            reason="Canonical authority applicability remains unresolved or generation-blocking",
                            owner="hub_authority_applicability", linked_type="trust",
                            linked_id=source_id, bundle_type=bundle_type,
                            context_id=source_id,
                        ))
    deduped = {}
    for item in items:
        key = (item["source_owner"], item["blocker_id"])
        deduped.setdefault(key, item)
    return list(deduped.values())


def _canonical_task_36_identity(connection, *, intake_id, firm_id):
    if not intake_id:
        return None
    if "intake_followup_tasks" not in _tables(connection):
        return None
    columns = _columns(connection, "intake_followup_tasks")
    identity_column = "id" if "id" in columns else ("task_id" if "task_id" in columns else None)
    if not identity_column:
        return None
    candidates = (36,) if identity_column == "id" else ("36", "TASK-36")
    for candidate in candidates:
        clauses, params = [f"{identity_column}=?"], [candidate]
        if "intake_id" in columns and intake_id:
            clauses.append("intake_id=?"); params.append(intake_id)
        if "firm_id" in columns:
            clauses.append("firm_id=?"); params.append(firm_id)
        row = connection.execute(
            f"SELECT {identity_column} FROM intake_followup_tasks WHERE {' AND '.join(clauses)}",
            params,
        ).fetchone()
        if row:
            return str(row[0])
    return None


def _docx(title: str, status: str, sections: list[tuple[str, Any]]) -> bytes:
    document = Document()
    document.add_heading(title, 0)
    document.add_paragraph(status)
    for heading, value in sections:
        document.add_heading(heading, level=1)
        if isinstance(value, (dict, list, tuple)):
            document.add_paragraph(json.dumps(value, indent=2, ensure_ascii=False, default=str))
        else:
            document.add_paragraph(str(value or "UNRESOLVED"))
    output = io.BytesIO(); document.save(output)
    # python-docx uses current ZIP member timestamps. Repack with fixed metadata
    # so identical controlled inputs produce identical content hashes.
    source = zipfile.ZipFile(io.BytesIO(output.getvalue()), "r")
    stable = io.BytesIO()
    with source, zipfile.ZipFile(stable, "w", zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    return stable.getvalue()


def _zip_bytes(files: dict[str, bytes]) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name in sorted(files):
            info = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            archive.writestr(info, files[name])
    return output.getvalue()


def _prior_bundle(connection, firm_id, source_id, template_id):
    connection.row_factory = sqlite3.Row
    row = connection.execute(
        """SELECT document_id FROM generated_documents
           WHERE firm_id=? AND source_record_id=? AND template_id=?
           ORDER BY created_at DESC, rowid DESC LIMIT 1""",
        (firm_id, source_id, template_id),
    ).fetchone()
    return row["document_id"] if row else None


def _lineage_terminals(rows, id_key, prior_key):
    predecessor_ids = {row.get(prior_key) for row in rows if row.get(prior_key)}
    return [row for row in rows if row[id_key] not in predecessor_ids]


def _resolve_will_content(connection, *, source_id, matter_id, firm_id):
    required = {"document_template_clause_revisions", "document_clause_generation_events",
                "document_clause_revision_content"}
    if not required.issubset(_tables(connection)):
        return {"clauses": [], "records": [], "revision_ids": [], "bindings": [],
                "conflicts": [], "missing": ["WILL_CLAUSE_CONTENT_OWNER"]}
    connection.row_factory = sqlite3.Row
    conditions, params = ["e.firm_id=?"], [firm_id]
    scope = ["e.intake_id=?", "(e.context_type='OTHER' AND e.context_id=?)"]
    params.extend([source_id, source_id])
    if matter_id:
        scope.extend(["e.matter_id=?", "(e.context_type='MATTER' AND e.context_id=?)"])
        params.extend([matter_id, matter_id])
    rows = [dict(r) for r in connection.execute(f"""SELECT e.*,r.clause_id,r.clause_key,
      r.revision_state,r.content_sha256,r.source_reference_id,r.prior_clause_revision_id,
      c.renderable_content,c.content_format
      FROM document_clause_generation_events e
      JOIN document_template_clause_revisions r ON r.clause_revision_id=e.clause_revision_id
      LEFT JOIN document_clause_revision_content c ON c.clause_revision_id=r.clause_revision_id
      WHERE {' AND '.join(conditions)} AND ({' OR '.join(scope)})""", params)]
    by_clause = {}
    for row in rows: by_clause.setdefault((row["template_id"], row["clause_id"]), []).append(row)
    clauses, records, revisions, bindings, conflicts, missing = [], [], [], [], [], []
    for (template_id, clause_id), candidates in by_clause.items():
        revision_rows = [dict(r) for r in connection.execute("""SELECT r.*,c.renderable_content,c.content_format
          FROM document_template_clause_revisions r
          LEFT JOIN document_clause_revision_content c ON c.clause_revision_id=r.clause_revision_id
          WHERE r.template_id=? AND r.clause_id=?""", (template_id, clause_id))]
        revision_terminals = _lineage_terminals(revision_rows, "clause_revision_id", "prior_clause_revision_id")
        viable = []
        for revision in revision_terminals:
            events = [r for r in candidates if r["clause_revision_id"] == revision["clause_revision_id"]]
            event_terminals = _lineage_terminals(events, "clause_generation_event_id", "prior_clause_generation_event_id")
            viable.extend([{**event, **revision} for event in event_terminals
                           if revision["revision_state"] == "APPROVED"
                           and event["generation_state"] in {"AUTHORIZED", "APPLIED"}
                           and int(event["human_confirmed"] or 0) == 1])
        if len(viable) > 1:
            conflicts.append({"clause_id": clause_id,
                              "clause_generation_event_ids": [r["clause_generation_event_id"] for r in viable]})
            continue
        if not viable:
            if any(r["revision_state"] == "APPROVED" for r in revision_terminals):
                missing.append({"clause_id": clause_id, "binding": "AUTHORIZED_GENERATION_EVENT"})
            continue
        row = viable[0]
        if row["renderable_content"] is None:
            missing.append({"clause_id": clause_id, "clause_revision_id": row["clause_revision_id"]})
            continue
        actual = _sha(row["renderable_content"].encode("utf-8"))
        if actual != row["content_sha256"]:
            raise ReviewBundleError("CLAUSE_CONTENT_HASH_MISMATCH")
        clauses.append({"clause_id": clause_id, "clause_key": row["clause_key"],
                        "text": row["renderable_content"]})
        revisions.append(row["clause_revision_id"])
        bindings.append({"clause_revision_id": row["clause_revision_id"],
                         "content_sha256": row["content_sha256"],
                         "content_format": row["content_format"]})
        records.append({"clause_revision_id": row["clause_revision_id"],
                        "clause_id": clause_id, "clause_key": row["clause_key"],
                        "generation_event_id": row["clause_generation_event_id"],
                        "source_reference_id": row["source_reference_id"],
                        "revision_lineage": [r["clause_revision_id"] for r in revision_rows],
                        "generation_provenance": row["provenance"]})
    return {"clauses": clauses, "records": records, "revision_ids": revisions,
            "bindings": bindings, "conflicts": conflicts, "missing": missing}


def _resolve_tr001_evidence(db_path, *, source_id, firm_id):
    from services.services_document_evidence import resolve_document_evidence
    try:
        result = resolve_document_evidence(db_path, firm_id=firm_id,
                                           context_type="TRUST", context_id=source_id)
    except sqlite3.DatabaseError:
        return {"relationships": [], "lineage": [], "conflicts": [],
                "missing": ["DOCUMENT_EVIDENCE_RELATIONSHIP_OWNER"]}
    required = {"GOVERNING_PRE_SIGNING_MASTER", "BENEFICIARY_SCHEDULE",
                "INITIAL_TRUSTEE_ACCEPTANCE", "SCHEDULE_A", "CLOSING_CONTROL_SHEET",
                "HISTORICAL_PROFESSIONAL_REVIEW_REFERENCE"}
    present = {row["evidence_role"] for row in result["relationships"]}
    result["missing"] = sorted(required - present)
    result["missing"].extend(
        {"evidence_role": row["evidence_role"], "binding": "DOCUMENT_INTRINSIC_INTEGRITY"}
        for row in result["relationships"]
        if not row.get("verified_content_sha256") or not row.get("source_revision_id")
    )
    return result


def generate_review_bundle(
    db_path, output_root, *, bundle_type, source_id, firm_id, owner_id,
    generated_by, intake_id=None, matter_id=None, governed_clauses=None,
    source_references=None, source_revision_ids=None, generation_batch_id=None,
    historical_reference_provenance=None, generated_at=None, bundle_id=None,
):
    """Generate and persist one review-only bundle from canonical records."""
    if bundle_type not in PROFILES:
        raise ReviewBundleError("unsupported_review_bundle_type")
    required = (source_id, firm_id, owner_id, generated_by)
    if not all(str(value or "").strip() for value in required):
        raise ReviewBundleError("bundle_scope_and_actor_required")
    profile = PROFILES[bundle_type]
    generated_at = generated_at or datetime.now(timezone.utc).isoformat()
    batch_id = generation_batch_id or "RGB-" + uuid.uuid4().hex[:12].upper()
    bundle_id = bundle_id or "RGB-" + uuid.uuid4().hex[:12].upper()
    clauses = []
    references = list(source_references or [])
    revisions = list(source_revision_ids or [])

    connection = sqlite3.connect(str(db_path)); connection.row_factory = sqlite3.Row
    try:
        needed = {"generated_documents", "professional_review_issues"}
        if not needed.issubset(_tables(connection)):
            raise ReviewBundleError("canonical_generated_document_or_review_owner_unavailable")
        source = _canonical_source(connection, bundle_type, source_id, firm_id)
        resolved_intake = intake_id or (source_id if bundle_type == WILL_BUNDLE else None)
        issues = _issues(connection, resolved_intake, firm_id)
        blockers = [row["issue_id"] for row in issues if _active_review_issue(row)]
        prior_id = _prior_bundle(connection, firm_id, source_id, profile["template_id"])
        issue_links = [{key: row.get(key) for key in (
            "issue_id", "linked_record_type", "linked_record_id", "status", "disposition",
            "reviewer_notes", "resolved_by", "resolved_capacity", "resolved_at",
        )} for row in issues]
        jurisdiction = source.get("jurisdiction") or ("New Jersey" if bundle_type == TR001_BUNDLE else "UNRESOLVED")
        will_resolution = (_resolve_will_content(connection, source_id=source_id,
                           matter_id=matter_id, firm_id=firm_id)
                           if bundle_type == WILL_BUNDLE else None)
        tr_resolution = (_resolve_tr001_evidence(db_path, source_id=source_id,
                         firm_id=firm_id) if bundle_type == TR001_BUNDLE else None)
        resolution = will_resolution or tr_resolution
        conflicts = resolution["conflicts"]
        missing = resolution["missing"]
        if conflicts: resolution_status = "CANONICAL_SOURCE_CONFLICT"
        elif missing: resolution_status = "MISSING_CANONICAL_BINDING"
        else: resolution_status = "CANONICAL_CONTENT_RESOLVED_WITH_INTENTIONAL_PLACEHOLDERS"
        if will_resolution:
            clauses = will_resolution["clauses"]
            revisions = list(dict.fromkeys(revisions + will_resolution["revision_ids"]))
            resolved_records = will_resolution["records"]
            content_bindings = will_resolution["bindings"]
            evidence_relationships = []
        else:
            resolved_records = tr_resolution["relationships"]
            content_bindings = []
            evidence_relationships = tr_resolution["relationships"]
            revisions = list(dict.fromkeys(revisions + [r["source_revision_id"] for r in evidence_relationships if r.get("source_revision_id")]))
        intentional = _derive_unresolved_items(
            connection, bundle_type=bundle_type, source=source,
            source_id=source_id, issues=issues,
        )
        task_36_identity = _canonical_task_36_identity(
            connection, intake_id=resolved_intake, firm_id=firm_id,
        )
        task_36_linked = task_36_identity is not None and any(
            str(row.get("linked_record_type") or "").strip() == "intake_followup_task"
            and str(row.get("linked_record_id") or "").strip() == task_36_identity
            for row in issues
        )
        if bundle_type == TR001_BUNDLE and task_36_identity is not None and not task_36_linked:
            missing = list(missing) + ["TASK_36_PROFESSIONAL_REVIEW_LINK"]
            if resolution_status != "CANONICAL_SOURCE_CONFLICT":
                resolution_status = "MISSING_CANONICAL_BINDING"
        identity = {
            "bundle_id": bundle_id, "candidate_id": bundle_id, "bundle_type": bundle_type,
            "matter_id": matter_id, "intake_id": resolved_intake,
            "trust_id": source_id if bundle_type == TR001_BUNDLE else None,
            "generated_at": generated_at, "generated_by": generated_by,
            "generation_batch_id": batch_id, "source_revision_ids": revisions,
            "source_snapshot_identifiers": {
                key: value for key, value in source.items()
                if value not in (None, "") and any(token in key.lower() for token in ("revision", "version", "snapshot"))
            },
            "controlling_canonical_source_references": references or [{"record_type": profile["source_type"], "record_id": source_id}],
            "professional_review_issues": issue_links, "unresolved_blockers": blockers,
            "jurisdiction_applicability_state": jurisdiction,
            "lifecycle_review_classification": profile["status"],
            "execution_status": "NON-EXECUTION", "supersedes": prior_id,
            "superseded_by": None, "historical_reference_provenance": list(historical_reference_provenance or []),
            "trust_status": TR001_TRUST_STATUS if bundle_type == TR001_BUNDLE else None,
            "canonical_content_resolution_status": resolution_status,
            "resolved_canonical_source_records": resolved_records,
            "resolved_revision_ids": revisions,
            "resolved_clause_content_bindings": content_bindings,
            "resolved_evidence_relationships": evidence_relationships,
            "intentional_unresolved_items": intentional,
            "missing_canonical_bindings": missing,
            "canonical_source_conflicts": conflicts,
        }
        common = [("Control identity", identity), ("Canonical source snapshot", source),
                  ("Professional Review issue links", issue_links), ("Unresolved blockers", blockers or ["NONE RECORDED"])]
        if bundle_type == WILL_BUNDLE:
            content = clauses or [{"state": "UNRESOLVED", "note": "No governed Will clause content was supplied; no provision was fabricated."}]
            components = {
                profile["components"][0]: _docx("Integrated Will — Professional Review Candidate", profile["status"], common + [("Governed clause content", content)]),
                profile["components"][1]: _docx("Professional Review Issues Memorandum", profile["status"], common),
                profile["components"][2]: _docx("Architecture & Decision Crosswalk", profile["status"], common + [("Revision provenance", revisions)]),
            }
        else:
            components = {
                profile["components"][0]: _docx("New Jersey Counsel Review Packet", profile["status"], common),
                profile["components"][1]: _docx("Tax Professional Review Packet", profile["status"], common),
                profile["components"][2]: _docx("Evidence & Provenance Index", profile["status"], common + [("Historical/reference provenance", identity["historical_reference_provenance"])]),
            }
        identity["component_filenames"] = list(profile["components"]) + ["component_sha256_manifest.sha256"]
        identity["component_sha256"] = {name: _sha(data) for name, data in components.items()}
        control_name = profile["components"][3]
        components[control_name] = (json.dumps(identity, indent=2, sort_keys=True, ensure_ascii=False) + "\n").encode("utf-8")
        component_hashes = {name: _sha(data) for name, data in components.items()}
        checksums = "".join(f"{digest}  {name}\n" for name, digest in sorted(component_hashes.items()))
        components["component_sha256_manifest.sha256"] = checksums.encode("utf-8")
        manifest_sha = _sha(components[control_name])
        package_bytes = _zip_bytes(components)
        package_sha = _sha(package_bytes)
        output_dir = Path(output_root) / bundle_id
        output_dir.mkdir(parents=True, exist_ok=False)
        package_path = output_dir / f"{bundle_id}.zip"
        package_path.write_bytes(package_bytes)
        metadata = {**identity, "component_sha256": component_hashes,
                    "manifest_sha256": manifest_sha, "package_sha256": package_sha,
                    "package_path": str(package_path)}
        columns = _columns(connection, "generated_documents")
        required_columns = {"document_id", "workspace_id", "trust_id", "template_id", "title", "content", "status", "created_by", "owner_id", "firm_id", "source_record_type", "source_record_id", "generation_basis"}
        if not required_columns.issubset(columns):
            raise ReviewBundleError("generated_documents_attribution_schema_unavailable")
        connection.execute("""INSERT INTO generated_documents
          (document_id,workspace_id,trust_id,template_id,title,content,status,created_by,owner_id,
           firm_id,source_record_type,source_record_id,generation_basis,created_at,updated_at)
          VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", (
            bundle_id, "professional-review", source_id if bundle_type == TR001_BUNDLE else None,
            profile["template_id"], f"{bundle_type} {bundle_id}", json.dumps(metadata, sort_keys=True),
            "PROFESSIONAL_REVIEW_NON_EXECUTION", generated_by, owner_id, firm_id,
            profile["source_type"], source_id, "CANONICAL_RECORD_PROFESSIONAL_REVIEW_BUNDLE",
            generated_at, generated_at,
        ))
        connection.commit()
        return metadata
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def get_review_bundle(db_path, *, bundle_id, firm_id, owner_id):
    connection = sqlite3.connect(str(db_path)); connection.row_factory = sqlite3.Row
    try:
        row = connection.execute(
            "SELECT content FROM generated_documents WHERE document_id=? AND firm_id=? AND owner_id=? AND workspace_id='professional-review'",
            (bundle_id, firm_id, owner_id),
        ).fetchone()
        return json.loads(row["content"]) if row else None
    finally:
        connection.close()
